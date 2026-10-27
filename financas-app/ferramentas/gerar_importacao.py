"""Transforma a Planilha financeira 2.0 (.xlsx) no arquivo que o app importa (darzan-dados.json).

Uso: python3 financas-app/ferramentas/gerar_importacao.py planilha.xlsx saida.json [pasta-de-fotos]

- Confere cada valor com o tipo da coluna no Supabase (darzan-producoes/financas/js/tabelas.js) e avisa o que não encaixar.
- Linhas sem ID ganham um (prefixo da tabela + número da linha), sempre o mesmo para a mesma linha.
- Fotos: se passar a pasta, equipamentos sem foto recebem a foto padrão pelo nome (zve10.jpg, iphone.jpg...).
- O arquivo gerado tem dados pessoais: não coloque no GitHub.
"""
import base64, datetime, io, json, os, re, subprocess, sys

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FOTOS = [(r'zv-?e?10.*(23|lente)|zve10 \+|c[aâ]mera', 'zve10'), (r'iphone', 'iphone'), (r'gimbal|scorp', 'scorp'), (r'ssd|xs1000', 'ssd'),
         (r'sokani|bast[aã]o', 'sokani'), (r'cart[aã]o de mem', 'cartao'), (r'bateria', 'bateria'), (r'mem[oó]ria ram|\bram\b', 'ram'), (r'case hd|hd externo', 'casehd')]


def tabelas():
    js = os.path.join(RAIZ, 'darzan-producoes/financas/js/tabelas.js')
    out = subprocess.run(['node', '-e', f'require({json.dumps(js)}); console.log(JSON.stringify(globalThis.TABELAS))'], capture_output=True, text=True, check=True)
    return json.loads(out.stdout)


def foto(caminho):
    from PIL import Image
    with Image.open(caminho) as im:
        im = im.convert('RGB')
        im.thumbnail((360, 360))
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=72, optimize=True)
    return 'data:image/jpeg;base64,' + base64.b64encode(buf.getvalue()).decode()


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    T = tabelas()
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    pasta = sys.argv[3] if len(sys.argv) > 3 else None
    avisos, abas = [], {}
    for aba, t in T.items():
        if aba not in wb.sheetnames:
            avisos.append(f'{aba}: a aba não existe na planilha')
            continue
        tipos = {nome: tipo for nome, _, tipo in t['colunas']}
        linhas = list(wb[aba].iter_rows(values_only=True))
        cab = [str(c).strip() if c is not None else '' for c in linhas[0]]
        regs = []
        for i, r in enumerate(linhas[1:], start=2):
            if not any(v not in (None, '', False) for v in r):
                continue
            reg = {}
            for c, v in zip(cab, r):
                if c != 'ID' and c not in tipos:
                    continue
                tipo = tipos.get(c, 'text')
                if v is None:
                    v = False if tipo == 'boolean' else ''
                elif c == 'PARCELA' and isinstance(v, datetime.datetime):
                    v = f'{v.day}/{v.month}'  # "2/9" que o Planilhas virou data
                elif tipo == 'date':
                    if isinstance(v, (datetime.datetime, datetime.date)):
                        v = v.strftime('%Y-%m-%d')
                    else:
                        avisos.append(f'{aba} linha {i}: "{c}" não é data ({v!r}); deixei vazio')
                        v = ''
                elif tipo == 'numeric':
                    if isinstance(v, bool) or not isinstance(v, (int, float)):
                        try:
                            v = float(str(v).replace('R$', '').replace('.', '').replace(',', '.').strip())
                        except ValueError:
                            avisos.append(f'{aba} linha {i}: "{c}" não é número ({v!r}); deixei vazio')
                            v = ''
                    if v != '' and float(v).is_integer():
                        v = int(v)
                elif tipo == 'boolean':
                    v = v is True or str(v).upper() == 'TRUE'
                else:
                    v = v.strftime('%d/%m/%Y') if isinstance(v, datetime.datetime) else str(v).strip()
                reg[c] = v
            reg['ID'] = str(reg.get('ID') or '').strip() or f"{t['prefixo']}{i:03d}"
            regs.append(reg)
        ids = [x['ID'] for x in regs]
        assert len(ids) == len(set(ids)), f'{aba}: IDs repetidos'
        abas[aba] = regs
    if pasta and os.path.isdir(pasta):
        for e in abas.get('EQUIPAMENTOS', []):
            if e.get('FOTO'):
                continue
            k = next((k for rx, k in FOTOS if re.search(rx, str(e.get('NOME', '')), re.I)), None)
            if k and os.path.exists(os.path.join(pasta, k + '.jpg')):
                e['FOTO'] = foto(os.path.join(pasta, k + '.jpg'))
    dados = {'app': 'Darzan Finanças', 'versao': 1, 'feito_em': datetime.datetime.now().isoformat(timespec='seconds'), 'abas': abas}
    with open(sys.argv[2], 'w') as f:
        json.dump(dados, f, ensure_ascii=False)
    print('ok:', sys.argv[2], f'({os.path.getsize(sys.argv[2]) // 1024} KB)', {k: len(v) for k, v in abas.items()})
    for a in avisos:
        print('  aviso:', a)


if __name__ == '__main__':
    main()
