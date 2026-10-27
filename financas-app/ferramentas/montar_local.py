"""Monta o app para abrir no computador (dist/local.html) e o arquivo de fotos para colar no Apps Script.

Uso:
  python3 ferramentas/montar_local.py planilha.xlsx [pasta-de-fotos] [foto-do-joao.jpg|.b64]

- planilha.xlsx: a Planilha financeira 2.0 baixada do Google (Arquivo → Fazer download → .xlsx).
- Gera dist/local.html (app + Google "de mentira" com uma cópia dos dados), dist/Fotos.html
  e dist/Index.html (tudo num arquivo só, para colar no Apps Script junto com o App.gs).
- dist/ não vai para o GitHub: tem dados pessoais.
"""
import base64, datetime, io, json, os, re, sys

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AS = os.path.join(RAIZ, 'apps-script')
DIST = os.path.join(RAIZ, 'dist')
ABAS = ['EXTRATO', 'CONTAS', 'CARTOES', 'PARCELADOS', 'FIXAS', 'EQUIPAMENTOS', 'TRABALHOS', 'TAREFAS', 'METAS', 'CATEGORIAS']


def celula(v):
    if isinstance(v, datetime.datetime) or isinstance(v, datetime.date):
        return {'$d': v.strftime('%Y-%m-%d')}
    return '' if v is None else v


def ler_planilha(arq):
    valores = openpyxl.load_workbook(arq, data_only=True)
    formulas = openpyxl.load_workbook(arq, data_only=False)
    abas, forms = {}, {}
    for nome in ABAS:
        if nome not in valores.sheetnames:
            continue
        grade = [[celula(v) for v in linha] for linha in valores[nome].iter_rows(values_only=True)]
        while grade and all(v == '' for v in grade[-1]):
            grade.pop()
        abas[nome] = grade
        f = {}
        for linha in formulas[nome].iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value.startswith('='):
                    f[f'{c.row},{c.column}'] = c.value  # o app só precisa saber que tem fórmula ali
        forms[nome] = f
    return {'abas': abas, 'formulas': forms}


def data_uri(caminho, lado=360, qualidade=72):
    from PIL import Image
    with Image.open(caminho) as im:
        im = im.convert('RGB')
        im.thumbnail((lado, lado))
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=qualidade, optimize=True)
    return 'data:image/jpeg;base64,' + base64.b64encode(buf.getvalue()).decode()


def fotos_html(pasta, foto_joao):
    partes = []
    if pasta and os.path.isdir(pasta):
        fotos = {os.path.splitext(n)[0]: data_uri(os.path.join(pasta, n)) for n in sorted(os.listdir(pasta)) if n.lower().endswith(('.jpg', '.jpeg', '.png', '.webp'))}
        partes.append('window.FOTOS_PADRAO = ' + json.dumps(fotos) + ';')
    if foto_joao and os.path.exists(foto_joao):
        if foto_joao.endswith('.b64'):
            uri = 'data:image/jpeg;base64,' + open(foto_joao).read().strip()
        else:
            uri = data_uri(foto_joao, 200)
        partes.append('window.FOTO_JOAO = ' + json.dumps(uri) + ';')
    return '<!-- Fotos do app (pessoal: não vai para o GitHub). -->\n<script>\n' + '\n'.join(partes) + '\n</script>\n'


def montar(index, incluir):
    return re.sub(r"<\?!= include\('(\w+)'\) \?>", lambda m: incluir(m.group(1)), index)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    os.makedirs(DIST, exist_ok=True)
    dados = ler_planilha(sys.argv[1])
    fotos = fotos_html(sys.argv[2] if len(sys.argv) > 2 else None, sys.argv[3] if len(sys.argv) > 3 else None)
    open(os.path.join(DIST, 'Fotos.html'), 'w').write(fotos)

    app_gs = open(os.path.join(AS, 'App.gs')).read()
    mock = open(os.path.join(RAIZ, 'testes', 'mock-google.js')).read()
    servidor = ('<script>window.DADOS_PLANILHA = ' + json.dumps(dados, ensure_ascii=False).replace('</', '<\\/') + ';</script>\n'
                '<script>\n' + mock + '\n</script>\n'
                '<script>\nwindow.SERVIDOR = (function () {\n' + app_gs + '\nreturn { carregar, lote };\n})();\n</script>\n')

    def incluir(nome):
        if nome == 'Fotos':
            return fotos
        txt = open(os.path.join(AS, nome + '.html')).read()
        return (servidor + txt) if nome == 'Base' else txt

    index = open(os.path.join(AS, 'Index.html')).read()
    # para colar no Apps Script: um arquivo só, com tudo dentro (sem o Google de mentira)
    pacote = montar(index, lambda nome: fotos if nome == 'Fotos' else open(os.path.join(AS, nome + '.html')).read())
    open(os.path.join(DIST, 'Index.html'), 'w').write(pacote)
    html = montar(index, incluir).replace('<base target="_top">', '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">')
    open(os.path.join(DIST, 'local.html'), 'w').write(html)
    print('ok:', os.path.join(DIST, 'local.html'), f'({len(html) // 1024} KB)', {k: len(v) - 1 for k, v in dados['abas'].items()})


if __name__ == '__main__':
    main()
