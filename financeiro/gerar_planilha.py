"""Gera a planilha financeira da Darzan Produções (versão 2).

Uso: python3 gerar_planilha.py saida.xlsx

As abas de lançamento têm o cabeçalho na linha 1 e um lançamento por linha.
O painel (financeiro/apps-script) lê essas abas direto do Google Planilhas,
procurando as colunas pelo nome do cabeçalho.
"""
import sys
from datetime import date, time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
N = 1000                      # linhas de lançamento preparadas em cada aba
FIRST, LAST = 2, N + 1

# ---------- estilos ----------
F = 'Arial'
DARK, GOLD, LINE = '1C1713', 'D6A85F', 'D9D2C5'
f_title = Font(name=F, size=18, bold=True, color=DARK)
f_sub = Font(name=F, size=10, italic=True, color='7A6F62')
f_head = Font(name=F, size=10, bold=True, color=GOLD)
f_body = Font(name=F, size=10, color='222222')
f_auto = Font(name=F, size=10, color='555555')
f_bold = Font(name=F, size=10, bold=True, color='222222')
f_sec = Font(name=F, size=12, bold=True, color=DARK)
fill_head = PatternFill('solid', fgColor=DARK)
fill_auto = PatternFill('solid', fgColor='EFEBE4')
fill_input = PatternFill('solid', fgColor='FFF3C4')
fill_hi = PatternFill('solid', fgColor='F6E7C8')
thin = Side(style='thin', color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)

BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"-"'
DATE = 'dd/mm/yyyy'
PCT = '0.0%;-0.0%;"-"'
INT = '0;-0;"-"'

MESES = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho', 'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']

wb = Workbook()


def rng(sheet, col):
    return f"'{sheet}'!${col}${FIRST}:${col}${LAST}"


def data_tab(ws, cols, fmts, tab_color):
    """cols: (nome, largura, 'input'|'auto'). Cabeçalho na linha 1."""
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    ws.sheet_properties.tabColor = tab_color
    ws.row_dimensions[1].height = 34
    for i, (name, width, kind) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = box
        ws.column_dimensions[L(i)].width = width
    for r in range(FIRST, LAST + 1):
        for i, (name, width, kind) in enumerate(cols, 1):
            c = ws.cell(row=r, column=i)
            c.border = box
            c.font = f_auto if kind == 'auto' else f_body
            if kind == 'auto':
                c.fill = fill_auto
            if fmts.get(i):
                c.number_format = fmts[i]
                c.alignment = Alignment(vertical='center')
            else:
                c.alignment = left
    ws.auto_filter.ref = f'A1:{L(len(cols))}{LAST}'


def dv_list(ws, src, col, prompt=None):
    dv = DataValidation(type='list', formula1=src, allow_blank=True, showErrorMessage=False)
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = 'Dica', prompt, True
    ws.add_data_validation(dv)
    dv.add(f'{col}{FIRST}:{col}{LAST}')


def fill(ws, rows):
    for r, row in enumerate(rows, FIRST):
        for j, v in enumerate(row, 1):
            if v is not None:
                ws.cell(row=r, column=j, value=v)


def pill(ws, col, values):
    for val, bg, fg in values:
        ws.conditional_formatting.add(f'{col}{FIRST}:{col}{LAST}', CellIsRule(
            operator='equal', formula=[f'"{val}"'], fill=PatternFill('solid', fgColor=bg),
            font=Font(name=F, color=fg, bold=True)))


OK, WARN, BAD = ('D7F0DD', '1E7B34'), ('FFF0C2', '8A6100'), ('FBDADA', 'A12828')
EX = 'EXEMPLO'

# =========================================================
# LISTAS (opções dos menus)
# =========================================================
SERVICOS = ['Casamento', 'Aniversário', 'Edição', 'Vídeo institucional', 'Storymaker', 'Fotografia', 'Chá revelação', 'Outro']
listas = {
    'A': ('Tipos de serviço', SERVICOS),
    'B': ('Status do orçamento', ['Aguardando', 'Fechado', 'Recusado']),
    'C': ('Forma de pagamento', ['Pix', 'Dinheiro', 'Cartão de crédito', 'Cartão de débito', 'Parcelado', 'Boleto', 'Transferência']),
    'D': ('Categorias de despesa', ['Combustível / transporte', 'Alimentação', 'Aplicativos / software', 'Cursos', 'Anúncios / marketing', 'Manutenção', 'Acessórios', 'Impostos / taxas', 'Outro']),
    'E': ('Categorias de equipamento', ['Câmera', 'Celular', 'Lente', 'Estabilizador / gimbal', 'Áudio', 'Iluminação', 'Drone', 'Computador', 'Armazenamento', 'Outro']),
    'F': ('Como o cliente conheceu', ['Instagram', 'Indicação', 'Site', 'Google', 'WhatsApp', 'Evento', 'Outro']),
    'G': ('Situação do equipamento', ['Em uso', 'Vendido', 'Quebrado / perdido']),
    'H': ('Meses', ['Ano inteiro'] + MESES),
}
LIST_LAST = 40


def lref(col):
    return f"=Listas!${col}$2:${col}${LIST_LAST}"


# =========================================================
# COMO USAR
# =========================================================
ws = wb.active
ws.title = 'Como usar'
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = '7A6F62'
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 30
ws.column_dimensions['C'].width = 95
ws['B2'] = 'Controle Financeiro — Darzan Produções'; ws['B2'].font = f_title
ws['B3'] = 'João Darzan · Videomaker & Storymaker · Campinas e região'; ws['B3'].font = f_sub
guia = [
    ('NO DIA A DIA', None),
    ('1. Fechou um trabalho?', 'Uma linha na aba Trabalhos: data, cliente, tipo de serviço, o que foi fechado, valor e quanto já recebeu. Pronto.'),
    ('2. Passou um orçamento?', 'Uma linha na aba Orçamentos, mesmo que o cliente ainda não tenha respondido. Quando ele responder, troque o Status para Fechado ou Recusado.'),
    ('3. Cliente novo?', 'Cadastre na aba Clientes (nome, WhatsApp, Instagram, como te conheceu). Use o mesmo nome nas outras abas.'),
    ('4. Pagou um freela?', 'Na própria linha do trabalho: quem foi (coluna Freelancers) e quanto custou (coluna Custo com freelancers).'),
    ('5. Gasto do negócio?', 'Combustível, aplicativo, anúncio, curso: aba Despesas. Equipamento novo: aba Equipamentos.'),
    ('', None),
    ('REGRAS SIMPLES', None),
    ('Colunas', 'BRANCAS você preenche. CINZAS são automáticas, não digite nelas. Não renomeie os títulos das colunas: o painel encontra os dados por eles.'),
    ('O que foi fechado', 'Escreva as entregas combinadas com o cliente: "Teaser + filme completo + 10 reels", "Storymaker 6h", "Ensaio 50 fotos". Isso aparece no painel em cada tipo de serviço.'),
    ('Datas e valores', 'Datas como 14/03/2026. Valores só com números (2500 ou 2500,50); o R$ aparece sozinho.'),
    ('Lucro', 'Lucro de cada trabalho = valor fechado − freelancers − outros custos do trabalho. No resumo do mês e do ano, as despesas gerais também são descontadas.'),
    ('Exemplos', 'As linhas marcadas com EXEMPLO servem só para você ver tudo funcionando. Apague essas linhas antes de começar.'),
    ('', None),
    ('PAINEL AUTOMÁTICO', None),
    ('O que é', 'Um painel no estilo da Darzan que lê esta planilha sozinho. Mexeu aqui, o painel já mostra. Só você acessa, com o seu login do Google.'),
    ('Como instalar', 'Uma vez só: Extensões → Apps Script → cole os dois arquivos que o Claude te mandou → Implantar. O passo a passo completo está na conversa e no arquivo LEIA-ME.'),
    ('Como abrir', 'Pelo menu Darzan → Abrir painel (aqui na planilha) ou pelo link do painel no celular.'),
    ('', None),
    ('RESUMO', None),
    ('Aba Resumo', 'Mesmo sem o painel, a aba Resumo mostra o ano e o mês que você escolher (células amarelas), mês a mês e por tipo de serviço.'),
]
r = 5
for a, b in guia:
    ws.cell(row=r, column=2, value=a)
    if b is None:
        ws.cell(row=r, column=2).font = Font(name=F, size=11, bold=True, color=GOLD)
        if a:
            for c in (2, 3):
                ws.cell(row=r, column=c).fill = fill_head
    else:
        ws.cell(row=r, column=2).font = f_bold
        ws.cell(row=r, column=3, value=b).font = f_body
        ws.cell(row=r, column=3).alignment = left
        ws.row_dimensions[r].height = 32
    r += 1

# =========================================================
# TRABALHOS (o coração da planilha)
# =========================================================
T = 'Trabalhos'
tw = wb.create_sheet(T)
tcols = [('Data do evento', 13, 'input'), ('Horário', 9, 'input'), ('Cliente', 24, 'input'),
         ('Tipo de serviço', 18, 'input'), ('O que foi fechado', 38, 'input'), ('Local / cidade', 16, 'input'),
         ('Valor fechado', 14, 'input'), ('Valor recebido', 14, 'input'), ('Forma de pagamento', 15, 'input'),
         ('Freelancers', 20, 'input'), ('Custo com freelancers', 14, 'input'), ('Outros custos do trabalho', 14, 'input'),
         ('A receber', 13, 'auto'), ('Lucro', 13, 'auto'), ('Pagamento', 11, 'auto'), ('Observações', 30, 'input')]
data_tab(tw, tcols, {1: DATE, 2: 'hh:mm', 7: BRL, 8: BRL, 11: BRL, 12: BRL, 13: BRL, 14: BRL}, '3A2E24')
for r in range(FIRST, LAST + 1):
    tw[f'M{r}'] = f'=IF(G{r}="","",G{r}-N(H{r}))'
    tw[f'N{r}'] = f'=IF(G{r}="","",G{r}-N(K{r})-N(L{r}))'
    tw[f'O{r}'] = f'=IF(G{r}="","",IF(N(H{r})>=G{r},"Pago",IF(N(H{r})>0,"Parcial","Pendente")))'
dv_list(tw, lref('A'), 'D')
dv_list(tw, lref('C'), 'I')
dv_list(tw, f"='Clientes'!$A${FIRST}:$A${LAST}", 'C', 'Escolha um cliente da aba Clientes ou digite o nome.')
pill(tw, 'O', [('Pago', *OK), ('Parcial', *WARN), ('Pendente', *BAD)])
trabalhos_ex = [
    (date(2026, 1, 24), time(16, 0), 'Cliente Exemplo 1', 'Casamento', 'Teaser + filme completo + 10 reels', 'Campinas', 4800, 4800, 'Pix', 'Freela A (2ª câmera)', 500, 120),
    (date(2026, 2, 14), time(20, 0), 'Cliente Exemplo 2', 'Aniversário', 'Storymaker 4h + teaser', 'Valinhos', 1200, 1200, 'Pix', '', 0, 40),
    (date(2026, 3, 7), time(9, 0), 'Empresa Exemplo', 'Vídeo institucional', 'Vídeo institucional 2 min + 3 cortes para redes', 'Campinas', 3500, 3500, 'Transferência', 'Freela B (drone)', 400, 60),
    (date(2026, 3, 21), time(15, 0), 'Cliente Exemplo 3', 'Chá revelação', 'Storymaker + vídeo da revelação', 'Hortolândia', 900, 900, 'Pix', '', 0, 30),
    (date(2026, 4, 11), time(17, 0), 'Cliente Exemplo 4', 'Casamento', 'Filme completo + cerimônia na íntegra', 'Indaiatuba', 5200, 5200, 'Parcelado', 'Freela A (2ª câmera)', 500, 150),
    (date(2026, 5, 3), time(10, 0), 'Cliente Exemplo 5', 'Fotografia', 'Ensaio externo · 60 fotos editadas', 'Campinas', 750, 750, 'Pix', '', 0, 0),
    (date(2026, 5, 20), None, 'Empresa Exemplo', 'Edição', 'Edição de 8 reels', 'Remoto', 960, 960, 'Pix', '', 0, 0),
    (date(2026, 6, 13), time(19, 0), 'Cliente Exemplo 6', 'Casamento', 'Teaser + storymaker + 15 reels', 'Sumaré', 3900, 3900, 'Pix', 'Freela C (storymaker)', 350, 90),
    (date(2026, 7, 18), time(21, 0), 'Cliente Exemplo 7', 'Aniversário', 'Festa 15 anos · teaser + filme', 'Campinas', 2400, 2400, 'Cartão de crédito', 'Freela C (storymaker)', 300, 60),
    (date(2026, 8, 22), time(16, 30), 'Cliente Exemplo 8', 'Casamento', 'Teaser + filme completo + drone', 'Jaguariúna', 5600, 5600, 'Parcelado', 'Freela A, Freela B', 900, 180),
    (date(2026, 9, 12), time(14, 0), 'Cliente Exemplo 9', 'Storymaker', 'Storymaker 6h · 20 stories só para o casal', 'Campinas', 1100, 1100, 'Pix', '', 0, 40),
    (date(2026, 10, 3), time(9, 0), 'Empresa Exemplo 2', 'Vídeo institucional', 'Vídeo de lançamento + 5 reels', 'Paulínia', 2800, 1400, 'Transferência', '', 0, 50),
    (date(2026, 10, 25), time(15, 0), 'Cliente Exemplo 10', 'Chá revelação', 'Teaser + fotos da revelação', 'Valinhos', 1300, 650, 'Pix', '', 0, 0),
    (date(2026, 11, 21), time(17, 0), 'Cliente Exemplo 11', 'Casamento', 'Filme completo + teaser + storymaker', 'Campinas', 6200, 3100, 'Parcelado', 'Freela A (2ª câmera)', 500, 0),
    (date(2026, 12, 12), time(20, 0), 'Cliente Exemplo 12', 'Aniversário', 'Storymaker 4h', 'Americana', 1100, 0, 'Pix', '', 0, 0),
]
fill(tw, [list(t) + [None, None, None, EX] for t in trabalhos_ex])

# =========================================================
# ORÇAMENTOS
# =========================================================
O = 'Orçamentos'
ow = wb.create_sheet(O)
ocols = [('Data do envio', 13, 'input'), ('Cliente', 24, 'input'), ('Tipo de serviço', 18, 'input'),
         ('O que foi orçado', 36, 'input'), ('Data do evento', 13, 'input'), ('Valor orçado', 14, 'input'),
         ('Status', 12, 'input'), ('Motivo / observações', 32, 'input')]
data_tab(ow, ocols, {1: DATE, 5: DATE, 6: BRL}, '3A2E24')
dv_list(ow, lref('A'), 'C'); dv_list(ow, lref('B'), 'G')
dv_list(ow, f"='Clientes'!$A${FIRST}:$A${LAST}", 'B', 'Escolha um cliente da aba Clientes ou digite o nome.')
pill(ow, 'G', [('Fechado', *OK), ('Aguardando', *WARN), ('Recusado', *BAD)])
orc_ex = []
for t in trabalhos_ex:
    d = t[0]
    envio = date(d.year if d.month > 1 else d.year - 1, d.month - 1 if d.month > 1 else 12, min(d.day, 28))
    orc_ex.append([envio, t[2], t[3], t[4], d, t[6], 'Fechado', EX])
orc_ex += [
    [date(2026, 2, 2), 'Cliente Exemplo 13', 'Casamento', 'Teaser + filme completo', date(2026, 6, 6), 6500, 'Recusado', EX + ' · achou caro'],
    [date(2026, 4, 15), 'Cliente Exemplo 14', 'Aniversário', 'Storymaker 3h', date(2026, 5, 30), 900, 'Recusado', EX + ' · data ocupada'],
    [date(2026, 6, 2), 'Empresa Exemplo 3', 'Vídeo institucional', 'Vídeo institucional 1 min', date(2026, 7, 10), 2500, 'Recusado', EX + ' · fechou com outro'],
    [date(2026, 8, 5), 'Cliente Exemplo 15', 'Fotografia', 'Ensaio de gestante', date(2026, 9, 1), 650, 'Recusado', EX],
    [date(2026, 9, 28), 'Cliente Exemplo 16', 'Casamento', 'Filme completo + storymaker', date(2027, 3, 20), 5900, 'Aguardando', EX],
    [date(2026, 10, 6), 'Empresa Exemplo 4', 'Edição', 'Edição de 12 reels por mês', date(2026, 11, 1), 1400, 'Aguardando', EX],
]
orc_ex.sort(key=lambda x: x[0])
fill(ow, orc_ex)

# =========================================================
# CLIENTES
# =========================================================
C = 'Clientes'
cw = wb.create_sheet(C)
ccols = [('Nome', 24, 'input'), ('WhatsApp', 16, 'input'), ('Instagram', 18, 'input'), ('Cidade', 15, 'input'),
         ('Como conheceu', 15, 'input'), ('Trabalhos', 10, 'auto'), ('Total fechado', 15, 'auto'), ('Observações', 30, 'input')]
data_tab(cw, ccols, {6: INT, 7: BRL}, '3A2E24')
for r in range(FIRST, LAST + 1):
    cw[f'F{r}'] = f'=IF(A{r}="","",COUNTIF({rng(T, "C")},A{r}))'
    cw[f'G{r}'] = f'=IF(A{r}="","",SUMIFS({rng(T, "G")},{rng(T, "C")},A{r}))'
dv_list(cw, lref('F'), 'E')
origens = ['Instagram', 'Indicação', 'Instagram', 'Site', 'Indicação', 'Instagram', 'Google', 'Indicação', 'Instagram', 'Indicação', 'Instagram', 'Site', 'Indicação', 'WhatsApp']
vistos = []
for t in trabalhos_ex:
    if t[2] not in [v[0] for v in vistos]:
        vistos.append((t[2], t[5]))
fill(cw, [[n, f'(19) 90000-{i + 1:04d}', '', cidade, origens[i % len(origens)], None, None, EX]
          for i, (n, cidade) in enumerate(vistos)])

# =========================================================
# DESPESAS (gastos gerais do negócio)
# =========================================================
D = 'Despesas'
dw = wb.create_sheet(D)
dcols = [('Data', 13, 'input'), ('Categoria', 22, 'input'), ('Descrição', 32, 'input'), ('Valor', 14, 'input'),
         ('Forma de pagamento', 16, 'input')]
data_tab(dw, dcols, {1: DATE, 4: BRL}, '3A2E24')
dv_list(dw, lref('D'), 'B'); dv_list(dw, lref('C'), 'E')
desp = []
for m in range(1, 11):
    desp.append([date(2026, m, 5), 'Aplicativos / software', 'Assinatura app de edição (EXEMPLO)', 60, 'Cartão de crédito'])
desp += [
    [date(2026, 3, 10), 'Anúncios / marketing', 'Impulsionamento Instagram (EXEMPLO)', 200, 'Cartão de crédito'],
    [date(2026, 6, 18), 'Cursos', 'Curso de color grading (EXEMPLO)', 350, 'Pix'],
    [date(2026, 8, 2), 'Manutenção', 'Limpeza de sensor (EXEMPLO)', 180, 'Pix'],
    [date(2026, 9, 9), 'Anúncios / marketing', 'Impulsionamento Instagram (EXEMPLO)', 250, 'Cartão de crédito'],
]
desp.sort(key=lambda x: x[0])
fill(dw, desp)

# =========================================================
# EQUIPAMENTOS
# =========================================================
E = 'Equipamentos'
ew = wb.create_sheet(E)
ecols = [('Equipamento', 28, 'input'), ('Categoria', 20, 'input'), ('Data da compra', 13, 'input'), ('Valor pago', 14, 'input'),
         ('Forma de pagamento', 16, 'input'), ('Situação', 15, 'input'), ('Observações', 30, 'input')]
data_tab(ew, ecols, {3: DATE, 4: BRL}, '3A2E24')
dv_list(ew, lref('E'), 'B'); dv_list(ew, lref('C'), 'E'); dv_list(ew, lref('G'), 'F')
fill(ew, [
    ['Câmera (exemplo)', 'Câmera', date(2025, 9, 15), 9500, 'Parcelado', 'Em uso', EX],
    ['Celular (exemplo)', 'Celular', date(2025, 11, 10), 7000, 'Parcelado', 'Em uso', EX],
    ['Gimbal (exemplo)', 'Estabilizador / gimbal', date(2025, 11, 10), 900, 'Pix', 'Em uso', EX],
    ['Microfone sem fio (exemplo)', 'Áudio', date(2026, 2, 5), 650, 'Cartão de crédito', 'Em uso', EX],
    ['Drone (exemplo)', 'Drone', date(2026, 7, 1), 4200, 'Parcelado', 'Em uso', EX],
])

# =========================================================
# RESUMO (dentro da planilha, sem precisar do painel)
# =========================================================
sw = wb.create_sheet('Resumo', 1)
sw.sheet_view.showGridLines = False
sw.sheet_properties.tabColor = GOLD
for k, v in {'A': 3, 'B': 34, 'C': 17, 'D': 17, 'E': 3, 'F': 20, 'G': 15, 'H': 15, 'I': 15, 'J': 12, 'K': 12, 'L': 12}.items():
    sw.column_dimensions[k].width = v
sw['B1'] = 'Resumo — Darzan Produções'; sw['B1'].font = f_title; sw.row_dimensions[1].height = 30
sw['B2'] = 'Tudo automático. Escolha o ano e o mês nas células amarelas.'; sw['B2'].font = f_sub
sw['B4'] = 'Ano:'; sw['B5'] = 'Mês:'
for c in ('B4', 'B5'):
    sw[c].font = f_bold; sw[c].alignment = Alignment(horizontal='right')
sw['C4'] = 2026; sw['C5'] = 'Outubro'
for c in ('C4', 'C5'):
    sw[c].fill = fill_input; sw[c].font = Font(name=F, size=12, bold=True, color='0000FF'); sw[c].alignment = center; sw[c].border = box
dvm = DataValidation(type='list', formula1=lref('H'), allow_blank=False)
sw.add_data_validation(dvm); dvm.add('C5')
ANO = '$C$4'
MI = 'MATCH($C$5,Listas!$H$3:$H$14,0)'          # 1..12 (erro se "Ano inteiro")


def head_row(r, labels, start=2):
    for i, t in enumerate(labels):
        c = sw.cell(row=r, column=start + i, value=t)
        c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = box
    sw.row_dimensions[r].height = 28


def between(sheet, col, datecol, a, b, extra=''):
    return f'SUMIFS({rng(sheet, col)},{rng(sheet, datecol)},">="&{a},{rng(sheet, datecol)},"<"&{b}{extra})'


def count_between(sheet, datecol, a, b, extra=''):
    return f'COUNTIFS({rng(sheet, datecol)},">="&{a},{rng(sheet, datecol)},"<"&{b}{extra})'


Y0, Y1 = f'DATE({ANO},1,1)', f'DATE({ANO}+1,1,1)'
M0, M1 = f'DATE({ANO},IFERROR({MI},1),1)', f'DATE({ANO},IFERROR({MI},12)+1,1)'
head_row(7, ['Indicador', 'No ano', 'No mês escolhido'])
okO = f',{rng(O, "G")},"Fechado"'
kpis = [
    ('Faturamento (valor fechado)', lambda a, b: between(T, 'G', 'A', a, b), BRL),
    ('Já recebido', lambda a, b: between(T, 'H', 'A', a, b), BRL),
    ('A receber', None, BRL),
    ('Freelancers', lambda a, b: between(T, 'K', 'A', a, b), BRL),
    ('Outros custos dos trabalhos', lambda a, b: between(T, 'L', 'A', a, b), BRL),
    ('Despesas gerais', lambda a, b: between(D, 'D', 'A', a, b), BRL),
    ('LUCRO', None, BRL),
    ('Margem de lucro', None, PCT),
    ('Trabalhos', lambda a, b: count_between(T, 'A', a, b), INT),
    ('Valor médio por trabalho', None, BRL),
    ('Orçamentos passados', lambda a, b: count_between(O, 'A', a, b), INT),
    ('Orçamentos fechados', lambda a, b: count_between(O, 'A', a, b, okO), INT),
    ('Taxa de fechamento', None, PCT),
]
for i, (label, fn, fmt) in enumerate(kpis):
    r = 8 + i
    sw.cell(row=r, column=2, value=label)
    for col, (a, b) in (('C', (Y0, Y1)), ('D', (M0, M1))):
        if fn:
            sw[f'{col}{r}'] = '=' + fn(a, b)
        sw[f'{col}{r}'].number_format = fmt
    for c in (2, 3, 4):
        cell = sw.cell(row=r, column=c); cell.border = box; cell.font = f_body
for col in 'CD':
    sw[f'{col}10'] = f'={col}8-{col}9'
    sw[f'{col}14'] = f'={col}8-{col}11-{col}12-{col}13'
    sw[f'{col}15'] = f'=IF({col}8=0,0,{col}14/{col}8)'
    sw[f'{col}17'] = f'=IF({col}16=0,0,{col}8/{col}16)'
    sw[f'{col}20'] = f'=IF({col}18=0,0,{col}19/{col}18)'
for c in (2, 3, 4):
    sw.cell(row=14, column=c).font = Font(name=F, size=11, bold=True, color=DARK); sw.cell(row=14, column=c).fill = fill_hi
sw['D21'] = '=IF(ISNA(MATCH($C$5,Listas!$H$3:$H$14,0)),"Escolha um mês na célula C5","")'
sw['D21'].font = Font(name=F, size=9, italic=True, color='A12828')

# Mês a mês
MR = 7
head_row(MR, ['Mês', 'Faturamento', 'Custos', 'Lucro', 'Trabalhos', 'Orçamentos', 'Fechados'], start=6)
for m in range(1, 13):
    r = MR + m
    a, b = f'DATE({ANO},{m},1)', f'DATE({ANO},{m}+1,1)'
    sw[f'F{r}'] = MESES[m - 1]
    sw[f'G{r}'] = '=' + between(T, 'G', 'A', a, b)
    sw[f'H{r}'] = f'={between(T, "K", "A", a, b)}+{between(T, "L", "A", a, b)}+{between(D, "D", "A", a, b)}'
    sw[f'I{r}'] = f'=G{r}-H{r}'
    sw[f'J{r}'] = '=' + count_between(T, 'A', a, b)
    sw[f'K{r}'] = '=' + count_between(O, 'A', a, b)
    sw[f'L{r}'] = '=' + count_between(O, 'A', a, b, okO)
TR = MR + 13
sw[f'F{TR}'] = 'TOTAL DO ANO'
for col in 'GHIJKL':
    sw[f'{col}{TR}'] = f'=SUM({col}{MR + 1}:{col}{MR + 12})'
for r in range(MR + 1, TR + 1):
    for col in 'FGHIJKL':
        c = sw[f'{col}{r}']; c.border = box; c.font = f_bold if r == TR else f_body
        c.number_format = BRL if col in 'GHI' else INT
        if r == TR:
            c.fill = fill_hi

# Por tipo de serviço
PR = 25
sw[f'B{PR - 1}'] = 'Por tipo de serviço (ano escolhido)'; sw[f'B{PR - 1}'].font = f_sec
head_row(PR, ['Tipo de serviço', 'Trabalhos', 'Faturamento', '', 'Valor médio', 'Lucro', 'Orçamentos', 'Taxa fech.'])
sw.cell(row=PR, column=5).fill = PatternFill(fill_type=None)
for i in range(len(SERVICOS)):
    r = PR + 1 + i
    sw[f'B{r}'] = f'=Listas!$A${2 + i}'
    t = f',{rng(T, "D")},$B{r}'
    o = f',{rng(O, "C")},$B{r}'
    sw[f'C{r}'] = '=' + count_between(T, 'A', Y0, Y1, t)
    sw[f'D{r}'] = '=' + between(T, 'G', 'A', Y0, Y1, t)
    sw[f'F{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
    sw[f'G{r}'] = f'=D{r}-{between(T, "K", "A", Y0, Y1, t)}-{between(T, "L", "A", Y0, Y1, t)}'
    sw[f'H{r}'] = '=' + count_between(O, 'A', Y0, Y1, o)
    sw[f'I{r}'] = f'=IF(H{r}=0,0,{count_between(O, "A", Y0, Y1, o + okO)}/H{r})'
    for col, fmt in zip('BCDFGHI', [None, INT, BRL, BRL, BRL, INT, PCT]):
        c = sw[f'{col}{r}']; c.border = box; c.font = f_body
        if fmt:
            c.number_format = fmt

ch = BarChart(); ch.type = 'col'; ch.title = 'Faturamento e lucro por mês'
ch.add_data(Reference(sw, min_col=7, min_row=MR, max_row=MR + 12), titles_from_data=True)
ch.add_data(Reference(sw, min_col=9, min_row=MR, max_row=MR + 12), titles_from_data=True)
ch.set_categories(Reference(sw, min_col=6, min_row=MR + 1, max_row=MR + 12))
ch.height, ch.width = 8, 20
ch.series[0].graphicalProperties.solidFill = 'D6A85F'
ch.series[1].graphicalProperties.solidFill = '3A2E24'
sw.add_chart(ch, 'B36')

# =========================================================
# LISTAS
# =========================================================
lw = wb.create_sheet('Listas')
lw.sheet_view.showGridLines = False
lw.sheet_properties.tabColor = 'BBBBBB'
for col, (title, items) in listas.items():
    c = lw[f'{col}1']; c.value = title; c.font = f_head; c.fill = fill_head; c.alignment = center
    lw.column_dimensions[col].width = 24
    for i, it in enumerate(items):
        lw[f'{col}{2 + i}'] = it
    for r in range(2, LIST_LAST + 1):
        lw[f'{col}{r}'].border = box; lw[f'{col}{r}'].font = f_body
lw.row_dimensions[1].height = 30
lw['J1'] = 'Pode adicionar opções nas linhas vazias. Não renomeie "Fechado", "Aguardando" e "Recusado".'
lw['J1'].font = f_sub

wb.active = 1
wb.save(OUT)
print('ok', wb.sheetnames)
