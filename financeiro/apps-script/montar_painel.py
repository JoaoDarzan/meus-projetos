"""Monta o Painel.html final: coloca os dados de exemplo da planilha dentro do modelo.

Uso: python3 montar_painel.py planilha.xlsx Painel.html
(a planilha precisa ter as fórmulas calculadas, por exemplo depois de abrir no LibreOffice)
"""
import json
import sys
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

ABAS = ['Trabalhos', 'Orçamentos', 'Clientes', 'Despesas', 'Equipamentos']


def valor(v):
    if isinstance(v, datetime):
        return v.strftime('%H:%M') if v.year < 1901 else v.strftime('%Y-%m-%d')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, time):
        return v.strftime('%H:%M')
    return '' if v is None else v


def main(xlsx, saida):
    wb = load_workbook(xlsx, data_only=True)
    abas = {}
    for nome in ABAS:
        ws = wb[nome]
        linhas = list(ws.iter_rows(values_only=True))
        cab = [str(h).strip() if h is not None else '' for h in linhas[0]]
        dados = []
        for row in linhas[1:]:
            if all(v in (None, '') for v in row):
                continue
            dados.append({h: valor(v) for h, v in zip(cab, row) if h})
        abas[nome] = dados
    demo = {'planilha': '', 'atualizado': '', 'abas': abas}
    modelo = (Path(__file__).parent / 'painel-modelo.html').read_text(encoding='utf-8')
    final = modelo.replace('/*__DEMO__*/null', json.dumps(demo, ensure_ascii=False))
    Path(saida).write_text(final, encoding='utf-8')
    print('ok', {k: len(v) for k, v in abas.items()}, len(final), 'bytes')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
